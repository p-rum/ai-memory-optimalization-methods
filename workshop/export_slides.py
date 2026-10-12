"""Optional PDF and editable PowerPoint exports; pip install python-pptx reportlab matplotlib."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import textwrap
R=Path(__file__).parent;slides=json.loads((R/'slides.json').read_text());D=json.loads((R/'results.json').read_text())
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
fig,axs=plt.subplots(1,2,figsize=(11,3.5),layout='constrained')
labels=[x['method'] for x in D['summary']];colors=['#8da1b8','#4da4ee','#f2b65a','#3dc6aa']
for ax,field,title in zip(axs,['complete','context_bytes'],['Úplné sady důkazů (z 24)','Průměrný kontext (UTF-8 bajty)']):
 vals=[x[field] for x in D['summary']];ax.bar(labels,vals,color=colors);ax.set_title(title);ax.spines[['top','right']].set_visible(False)
 for i,v in enumerate(vals):ax.text(i,v+max(vals)*.02,str(v),ha='center')
 ax.set_ylim(0,max(vals)*1.18)
fig.savefig(R/'results-chart.png',dpi=180);plt.close(fig)
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
for i,(title,sub,bullets,note) in enumerate(slides):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('101728')
 def text(x,y,w,h,txt,size,color='F0F4FF'):
  box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.word_wrap=True
  for j,line in enumerate(txt.split('\n')):
   p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color);p.font.name='DejaVu Sans';p.space_after=Pt(14)
 text(.6,.2,12,.4,f'LLM MEMORY LAB · {i+1:02}/12',12,'A9BBD6');text(.6,.8,12,1.1,title,32);text(.6,1.95,12,.6,sub,21,'75DFC8')
 if i==7:slide.shapes.add_picture(str(R/'results-chart.png'),Inches(.8),Inches(2.8),width=Inches(11.6))
 else:text(.7,2.8,12,4.3,'\n'.join('• '+b for b in bullets),23)
 slide.notes_slide.notes_text_frame.text=note
prs.save(R/'presentation.pptx')
pdfmetrics.registerFont(TTFont('DejaVu','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
c=canvas.Canvas(str(R/'presentation.pdf'),pagesize=(960,540));c.setTitle('Paměť LLM agenta: 30minutový workshop')
for i,(title,sub,bullets,note) in enumerate(slides):
 c.setFillColorRGB(.063,.09,.157);c.rect(0,0,960,540,fill=1,stroke=0)
 c.setFillColorRGB(.67,.75,.84);c.setFont('DejaVu',11);c.drawString(42,507,f'LLM MEMORY LAB · {i+1:02}/12')
 c.setFillColorRGB(.94,.96,1);c.setFont('DejaVu',29);c.drawString(42,450,title)
 c.setFillColorRGB(.46,.87,.78);c.setFont('DejaVu',18);c.drawString(42,405,sub)
 if i==7:c.drawImage(str(R/'results-chart.png'),42,80,width=870,height=277)
 else:
  c.setFillColorRGB(.94,.96,1);c.setFont('DejaVu',19);y=352
  for b in bullets:
   for j,line in enumerate(textwrap.wrap(b,76)):
    c.drawString(42,y,('• ' if j==0 else '  ')+line);y-=28
   y-=14
 c.showPage()
c.save()
print('Created presentation.pptx, presentation.pdf, results-chart.png')
